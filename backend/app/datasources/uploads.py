"""Turns uploaded CSV / Excel files into a small SQLite database that DataPilot can query.

Each CSV file, and each non-empty Excel sheet, becomes one table. Column types are inferred from
the values (INTEGER, REAL, DATE, DATETIME, otherwise TEXT) so that sums, averages and date
functions work in the generated SQL. Names are normalised to lower_snake_case identifiers.
"""

import csv
import io
import os
import re
import sqlite3
import tempfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from app.core.errors import BadRequestError

MAX_TABLES = 20
MAX_COLUMNS = 200
MAX_ROWS_PER_TABLE = 500_000
CSV_EXTENSIONS = (".csv", ".tsv", ".txt")
EXCEL_EXTENSIONS = (".xlsx", ".xlsm")

_INT = re.compile(r"^[+-]?\d+$")
_FLOAT = re.compile(r"^[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?$")
_GROUPED = re.compile(r"^[+-]?\d{1,3}(,\d{3})+(\.\d+)?$")  # 1,234,567.89
_LEADING_ZERO = re.compile(r"^[+-]?0\d")  # 011001, 007: codes, not numbers
# Day-first before month-first: an ambiguous column like 03/04/2025 is read the Indian/UK way.
_DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%d %b %Y", "%d-%b-%Y", "%b %d, %Y")
_DATETIME_FORMATS = ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M", "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S")


@dataclass
class ParsedTable:
    name: str
    columns: list[str]
    rows: list[list[Any]]


def identifier(raw: str, fallback: str) -> str:
    name = re.sub(r"[^0-9a-zA-Z]+", "_", str(raw).strip()).strip("_").lower()
    if not name:
        name = fallback
    if name[0].isdigit():
        name = f"c_{name}"
    return name[:60]


def _dedupe(names: Iterable[str], taken: set[str] | None = None) -> list[str]:
    seen = set(taken or ())
    out = []
    for name in names:
        candidate, n = name, 2
        while candidate in seen:
            candidate, n = f"{name}_{n}", n + 1
        seen.add(candidate)
        out.append(candidate)
    return out


# -- reading ---------------------------------------------------------------------------------


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


def _read_csv(filename: str, content: bytes) -> Iterator[tuple[str, list[list[Any]]]]:
    text = _decode(content)
    if filename.lower().endswith(".tsv"):
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    csv.field_size_limit(10_000_000)
    yield Path(filename).stem, list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _read_excel(filename: str, content: bytes) -> Iterator[tuple[str, list[list[Any]]]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise BadRequestError(f"{filename} could not be read as an Excel workbook.") from exc
    try:
        sheets = [(sheet.title, [list(r) for r in sheet.iter_rows(values_only=True)]) for sheet in workbook.worksheets]
        sheets = [(title, rows) for title, rows in sheets if any(not _blank(v) for r in rows for v in r)]
        stem = Path(filename).stem
        for title, rows in sheets:
            yield (stem if len(sheets) == 1 else f"{stem}_{title}"), rows
    finally:
        workbook.close()


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _to_table(name: str, raw_rows: list[list[Any]], filename: str) -> ParsedTable | None:
    rows = [r for r in raw_rows if any(not _blank(v) for v in r)]
    if not rows:
        return None
    header, body = rows[0], rows[1:]
    width = max(len(r) for r in rows)
    # Drop trailing columns that are empty everywhere (common in exported spreadsheets).
    while width and all(len(r) < width or _blank(r[width - 1]) for r in rows):
        width -= 1
    if width == 0:
        return None
    if width > MAX_COLUMNS:
        raise BadRequestError(f"{filename} has {width} columns; the limit is {MAX_COLUMNS}.")
    if len(body) > MAX_ROWS_PER_TABLE:
        raise BadRequestError(f"{filename} has {len(body):,} rows; the limit is {MAX_ROWS_PER_TABLE:,} per table.")
    header = list(header) + [None] * (width - len(header))
    columns = _dedupe(identifier("" if _blank(h) else h, f"column_{i + 1}") for i, h in enumerate(header[:width]))
    body = [(list(r) + [None] * (width - len(r)))[:width] for r in body]
    return ParsedTable(name=name, columns=columns, rows=body)


def parse_files(files: list[tuple[str, bytes]]) -> list[ParsedTable]:
    tables: list[ParsedTable] = []
    for filename, content in files:
        lower = filename.lower()
        if lower.endswith(CSV_EXTENSIONS):
            sources = _read_csv(filename, content)
        elif lower.endswith(EXCEL_EXTENSIONS):
            sources = _read_excel(filename, content)
        elif lower.endswith(".xls"):
            raise BadRequestError(f"{filename} is an old .xls workbook. Save it as .xlsx or .csv and upload again.")
        else:
            raise BadRequestError(f"{filename} isn't a supported file. Upload .csv, .tsv or .xlsx files.")
        for name, raw_rows in sources:
            table = _to_table(name, raw_rows, filename)
            if table:
                tables.append(table)
    if not tables:
        raise BadRequestError("The uploaded files contain no data.")
    if len(tables) > MAX_TABLES:
        raise BadRequestError(f"That's {len(tables)} tables; the limit is {MAX_TABLES} per data source.")
    for table, name in zip(tables, _dedupe(identifier(t.name, "table") for t in tables), strict=True):
        table.name = name
    return tables


# -- typing ----------------------------------------------------------------------------------


def _parse_temporal(text: str, formats: tuple[str, ...]) -> str | None:
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).isoformat(sep=" ")
        except ValueError:
            continue
    return None


def _column_type(values: list[Any]) -> tuple[str, Any]:
    """Pick a SQL type for a column and return it with a converter for its values."""
    present = [v for v in values if not _blank(v)]
    if not present:
        return "TEXT", lambda v: None if _blank(v) else str(v)

    if all(isinstance(v, datetime | date) and not isinstance(v, time) for v in present):
        has_time = any(isinstance(v, datetime) and v.time() != time(0) for v in present)
        if has_time:
            return "DATETIME", lambda v: None if _blank(v) else v.isoformat(sep=" ")
        return "DATE", lambda v: None if _blank(v) else (v.date() if isinstance(v, datetime) else v).isoformat()
    if all(isinstance(v, bool) for v in present):
        return "INTEGER", lambda v: None if _blank(v) else int(v)
    if all(isinstance(v, int | float) and not isinstance(v, bool) for v in present):
        if all(float(v).is_integer() for v in present):
            return "INTEGER", lambda v: None if _blank(v) else int(v)
        return "REAL", lambda v: None if _blank(v) else float(v)

    texts = [str(v).strip() for v in present]
    # Leading zeros (PIN codes, IDs like 00123) stay text so they aren't mangled.
    numeric = not any(_LEADING_ZERO.match(t) for t in texts)
    if numeric and all(_INT.match(t) for t in texts):
        return "INTEGER", lambda v: None if _blank(v) else int(str(v).strip())
    if numeric and all(_FLOAT.match(t) or _GROUPED.match(t) for t in texts):
        return "REAL", lambda v: None if _blank(v) else float(str(v).strip().replace(",", ""))
    for fmt in _DATE_FORMATS:
        if all(_parse_temporal(t, (fmt,)) for t in texts):
            return (
                "DATE",
                lambda v, f=fmt: None if _blank(v) else datetime.strptime(str(v).strip(), f).date().isoformat(),
            )
    for fmt in _DATETIME_FORMATS:
        if all(_parse_temporal(t, (fmt,)) for t in texts):
            return (
                "DATETIME",
                lambda v, f=fmt: None if _blank(v) else datetime.strptime(str(v).strip(), f).isoformat(sep=" "),
            )
    return "TEXT", lambda v: None if _blank(v) else (v.isoformat() if isinstance(v, date | time) else str(v).strip())


# -- building --------------------------------------------------------------------------------


def _quote(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def build_sqlite(tables: list[ParsedTable]) -> bytes:
    """Write the tables into a fresh SQLite database and return its bytes."""
    fd, path = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    try:
        conn = sqlite3.connect(path)
        try:
            for t in tables:
                typed = [_column_type([row[i] for row in t.rows]) for i in range(len(t.columns))]
                cols = ", ".join(f"{_quote(c)} {sql_type}" for c, (sql_type, _) in zip(t.columns, typed, strict=True))
                conn.execute(f"CREATE TABLE {_quote(t.name)} ({cols})")
                placeholders = ", ".join("?" for _ in t.columns)
                converters = [convert for _, convert in typed]
                conn.executemany(
                    f"INSERT INTO {_quote(t.name)} VALUES ({placeholders})",
                    ([convert(v) for convert, v in zip(converters, row, strict=True)] for row in t.rows),
                )
            conn.commit()
            conn.execute("VACUUM")
        finally:
            conn.close()
        return Path(path).read_bytes()
    finally:
        Path(path).unlink(missing_ok=True)


def describe(tables: list[ParsedTable]) -> str:
    return ", ".join(f"{t.name} ({len(t.rows):,} rows)" for t in tables)
